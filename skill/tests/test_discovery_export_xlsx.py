#!/usr/bin/env python3
"""discover Excel(.xlsx) 导出回归（P2 四大区选品簿，对标上品帮 §10 方法论吸收）。

承诺：
① 结构：两行表头——第 1 行四大区合并单元格（基础信息/销售数据/尺寸重量/我的定价）、
   第 2 行中文列名；四大区列数合计 = CSV 全字段数（只归区重排，不增不减）；
② 值语义与 CSV 同源（ozon_discovery._candidate_row 单一样事实源），数据行数一致；
③ 原子写 + 占用重试：成功后无 .tmp- 残留；目标被占用（PermissionError）退避重试后
   成功；彻底失败报 RuntimeError 且旧文件不破坏；
④ cli --export 后缀路由：.xlsx（大小写不敏感）→ export_to_xlsx，其余仍走 CSV。

候选 fixture 用合成 6 条真实形状候选（asdict dict 同构，字段有值/零值/None 混排）——
不读 data/discovery/ 本机工件（gitignore，Docker/CI 无此文件，2026-09-10 CI 修）。

运行：
    cd skill && .venv314/bin/python -m pytest tests/test_discovery_export_xlsx.py -q
"""
from __future__ import annotations

import csv
import os
import sys
import zipfile
from dataclasses import fields as dataclass_fields
from pathlib import Path
from unittest import mock

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

_ZONE_TITLES = ["基础信息", "销售数据", "尺寸重量", "我的定价"]


def _fixture_candidates():
    """合成 6 条真实形状候选（走真实 dict→ProductCandidate 过滤通道，向前兼容）。"""
    from scripts.lib.ozon_discovery import ProductCandidate
    valid = {f.name for f in dataclass_fields(ProductCandidate)}
    raw = [
        {"ozon_product_id": "1005570349", "ozon_title": "Катунь Serving Board",
         "ozon_price": 953.0, "status": "profitable",
         "ozon_url": "https://www.ozon.ru/product/1005570349",
         "ozon_images": ["https://ir-20.ozonstatic.cn/main.jpg"],
         "match_1688_url": "https://detail.1688.com/offer/1043458962058.html",
         "match_1688_title": "卡吞餐盘", "match_1688_price": 18.5,
         "match_1688_category_name": "餐具", "session_count": 0,
         "days_in_promo": 12, "profit_margin": 22.5, "weight_g": 640},
        {"ozon_product_id": "1622910561", "ozon_title": "Термос для воды 1л",
         "ozon_price": 1580.0, "status": "profitable",
         "match_1688_price": 32.0, "session_count": 3,
         "profit_margin": 18.2, "weight_g": 480, "days_in_promo": 0},
        {"ozon_product_id": "881234501", "ozon_title": "Коврик для мыши",
         "ozon_price": 420.0, "status": "watch",
         "match_1688_price": 6.8, "session_count": 55,
         "profit_margin": 4.1, "weight_g": 120, "days_in_promo": 45},
        {"ozon_product_id": "9017753210", "ozon_title": "Органайзер для косметики",
         "ozon_price": 736.0, "status": "profitable",
         "match_1688_price": 12.4, "session_count": 12,
         "profit_margin": 31.7, "weight_g": 260, "days_in_promo": 7},
        {"ozon_product_id": "7021133445", "ozon_title": "Набор чашек керамика",
         "ozon_price": 1890.0, "status": "skip",
         "match_1688_price": 58.0, "session_count": 210,
         "profit_margin": -2.3, "weight_g": 1500, "days_in_promo": 30},
        {"ozon_product_id": "5590123478", "ozon_title": "Подставка для телефона",
         "ozon_price": 310.0, "status": "watch",
         "match_1688_price": 4.2, "session_count": 8,
         "profit_margin": 12.9, "weight_g": 90, "days_in_promo": 0},
    ]
    return [ProductCandidate(**{k: v for k, v in d.items() if k in valid}) for d in raw]


def _mk_candidate():
    """合成候选（有值的字段尽量多，供值级对齐断言）。"""
    from scripts.lib.ozon_discovery import ProductCandidate
    c = ProductCandidate(
        ozon_product_id="1005570349",
        ozon_title="Катунь Serving Board",
        ozon_price=953.0,
    )
    c.status = "profitable"
    c.ozon_url = "https://www.ozon.ru/product/1005570349"
    c.ozon_images = ["https://ir-20.ozonstatic.cn/main.jpg"]
    c.match_1688_url = "https://detail.1688.com/offer/1043458962058.html"
    c.match_1688_title = "卡吞餐盘"
    c.match_1688_price = 18.5
    c.match_1688_category_name = "餐具"
    c.session_count = 0            # 真实 0 保留
    c.days_in_promo = 12
    c.profit_margin = 22.5
    c.weight_g = 640
    return c


def _xlsx_header_keys():
    """按四区顺序拼出 xlsx 列键序（与实现共用同一 spec）。"""
    from scripts.lib.ozon_discovery import _EXPORT_XLSX_ZONES
    return [key for _, cols in _EXPORT_XLSX_ZONES for key, _ in cols]


# ── ① 结构：四大区两行表头 ────────────────────────────────────────────────

def test_zone_spec_covers_exactly_csv_fields():
    """四大区列数合计 = CSV 字段数；字段集一致、无重复、区序固定。"""
    from scripts.lib.ozon_discovery import _EXPORT_FIELDS, _EXPORT_XLSX_ZONES
    assert [z for z, _ in _EXPORT_XLSX_ZONES] == _ZONE_TITLES, "四大区名与区序"
    keys = _xlsx_header_keys()
    assert len(keys) == len(_EXPORT_FIELDS), "四区列合计必须等于 CSV 字段数"
    assert sorted(keys) == sorted(_EXPORT_FIELDS), "字段集一致（只归区重排）"
    assert len(set(keys)) == len(keys), "xlsx 列不得重复"


def test_xlsx_two_row_header_with_zone_merges(tmp_path):
    """两行表头：第 1 行四区合并单元格（区名+连续跨度），第 2 行中文列名。"""
    from openpyxl import load_workbook

    from scripts.lib.ozon_discovery import _EXPORT_XLSX_ZONES, export_to_xlsx
    cands = _fixture_candidates()
    out = tmp_path / "export.xlsx"
    export_to_xlsx(cands, str(out))
    assert out.exists(), "xlsx 文件已生成"

    ws = load_workbook(out).active
    assert ws.title == "选品候选"
    total = sum(len(cols) for _, cols in _EXPORT_XLSX_ZONES)
    assert ws.max_column == total
    assert ws.max_row == 2 + len(cands), "两行表头 + 数据行数 = 候选数"

    # 第 1 行合并区：4 块、行号 1、跨度与各区分列数一致、锚点值=区名
    merges = sorted(ws.merged_cells.ranges, key=lambda r: r.min_col)
    assert len(merges) == 4, f"四大区各一个合并单元格, got {len(merges)}"
    col = 1
    for (zone, cols), rng in zip(_EXPORT_XLSX_ZONES, merges):
        assert (rng.min_row, rng.max_row) == (1, 1)
        assert rng.min_col == col and rng.max_col == col + len(cols) - 1, f"{zone} 跨度连续"
        assert ws.cell(row=1, column=rng.min_col).value == zone
        col = rng.max_col + 1

    # 第 2 行：中文列名逐列对齐 spec
    for idx, (_key, label) in enumerate(
            [pair for _, cols in _EXPORT_XLSX_ZONES for pair in cols], start=1):
        assert ws.cell(row=2, column=idx).value == label, f"列 {idx} 中文列名"


def test_xlsx_empty_candidates_header_only(tmp_path):
    """空候选：仅表头两行，不报错。"""
    from openpyxl import load_workbook

    from scripts.lib.ozon_discovery import export_to_xlsx
    out = tmp_path / "empty.xlsx"
    export_to_xlsx([], str(out))
    ws = load_workbook(out).active
    assert ws.max_row == 2


# ── ② 值语义与 CSV 同源 ──────────────────────────────────────────────────

def test_xlsx_values_match_csv_row(tmp_path):
    """同一候选：xlsx 每格值（字符串化）与 CSV 同名列逐格一致。"""
    from openpyxl import load_workbook

    from scripts.lib.ozon_discovery import export_to_csv, export_to_xlsx
    c = _mk_candidate()
    xlsx_path = tmp_path / "a.xlsx"
    csv_path = tmp_path / "a.csv"
    export_to_xlsx([c], str(xlsx_path))
    export_to_csv([c], str(csv_path))

    def _norm(v):
        # openpyxl 读回数值不带尾零（953.0→953）：可数值化的按 float 归一比较
        s = "" if v is None else str(v)
        try:
            return repr(float(s))
        except ValueError:
            return s

    with open(csv_path, encoding="utf-8-sig") as f:
        csv_row = list(csv.DictReader(f))[0]
    ws = load_workbook(xlsx_path).active
    for idx, key in enumerate(_xlsx_header_keys(), start=1):
        cell = ws.cell(row=3, column=idx).value
        assert _norm(cell) == _norm(csv_row[key]), \
            f"{key}: xlsx={cell!r} != csv={csv_row[key]!r}"


def test_xlsx_fixture_data_rows(tmp_path):
    """真实形状 fixture（6 条）：数据行数=6，profitable 候选结论列透传。"""
    from openpyxl import load_workbook

    from scripts.lib.ozon_discovery import export_to_xlsx
    cands = _fixture_candidates()
    assert len(cands) == 6
    out = tmp_path / "fx.xlsx"
    export_to_xlsx(cands, str(out))
    ws = load_workbook(out).active
    assert ws.max_row == 2 + 6
    verdict_col = _xlsx_header_keys().index("verdict") + 1
    got = [ws.cell(row=r, column=verdict_col).value for r in range(3, 3 + 6)]
    assert got == [c.status for c in cands], "verdict 列=候选 status 透传"


# ── ③ 原子写 + 占用重试 ──────────────────────────────────────────────────

def test_xlsx_atomic_write_no_tmp_residue(tmp_path):
    """成功导出（含覆盖旧文件）后，同目录无 .tmp- 残留。"""
    from scripts.lib.ozon_discovery import export_to_xlsx
    out = tmp_path / "a.xlsx"
    export_to_xlsx([_mk_candidate()], str(out))
    export_to_xlsx([_mk_candidate()], str(out))   # 覆盖写也不残留
    assert list(tmp_path.glob(".tmp-*")) == []


def _patch_save(side_effect):
    return mock.patch("openpyxl.workbook.workbook.Workbook.save",
                      side_effect=side_effect)


def test_xlsx_retry_succeeds_after_occupied_target(tmp_path):
    """目标被 Excel 占用（PermissionError）：重试后成功，内容完整，无残留。"""
    from openpyxl import load_workbook
    from openpyxl.workbook.workbook import Workbook

    from scripts.lib.ozon_discovery import export_to_xlsx
    out = tmp_path / "busy.xlsx"
    calls = {"n": 0}
    real_save = Workbook.save

    def _flaky_save(wb_self, path):
        # 首次模拟占用；重试放行真实保存（否则 tmp 不存在，os.replace 必 ENOENT）
        calls["n"] += 1
        if calls["n"] == 1:
            raise PermissionError(13, "文件被占用（模拟 Excel 打开）")
        return real_save(wb_self, path)

    with mock.patch.object(Workbook, "save", _flaky_save):
        export_to_xlsx([_mk_candidate()], str(out))
    assert calls["n"] == 2, "第一次被占用，第二次重试成功"
    ws = load_workbook(out).active
    assert ws.max_row == 3
    assert list(tmp_path.glob(".tmp-*")) == []


def test_xlsx_failure_keeps_old_file_and_no_residue(tmp_path):
    """彻底失败：报 RuntimeError（中文明确报错），旧文件原样保留，无临时残留。"""
    from openpyxl import load_workbook

    from scripts.lib.ozon_discovery import export_to_xlsx
    out = tmp_path / "old.xlsx"
    export_to_xlsx([_mk_candidate()], str(out))   # 先有一份旧文件
    before = out.read_bytes()

    def _always_busy(path):
        raise PermissionError(13, "EBUSY（模拟 Excel 一直占用）")

    with _patch_save(_always_busy):
        with pytest.raises(RuntimeError) as exc_info:
            export_to_xlsx([_mk_candidate()], str(out))
    assert "占用" in str(exc_info.value), "报错须人话说明文件被占用"
    assert out.read_bytes() == before, "旧文件绝不被半写破坏"
    assert list(tmp_path.glob(".tmp-*")) == [], "失败也不残留临时文件"
    load_workbook(out)  # 旧文件仍可正常打开


# ── ④ cli --export 后缀路由 ──────────────────────────────────────────────

def test_route_xlsx_suffix_case_insensitive(tmp_path):
    """.xlsx/.XLSX 后缀 → export_to_xlsx（产物是 zip 容器的 xlsx）。"""
    from scripts.cli import _route_discovery_export
    for name in ("out.xlsx", "OUT.XLSX"):
        p = tmp_path / name
        ret = _route_discovery_export([_mk_candidate()], str(p))
        assert ret == str(p)
        assert zipfile.is_zipfile(str(p)), f"{name} 应走 Excel 导出"


def test_route_non_xlsx_still_csv(tmp_path):
    """非 .xlsx（无后缀/.csv）仍走 CSV 导出。"""
    from scripts.cli import _route_discovery_export
    p = tmp_path / "out.csv"
    ret = _route_discovery_export([_mk_candidate()], str(p))
    assert ret == str(p)
    with open(p, encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        assert reader.fieldnames and "product_id" in reader.fieldnames
        assert len(list(reader)) == 1

    p2 = tmp_path / "noext"           # 无后缀 → CSV（旧行为兜底）
    ret2 = _route_discovery_export([_mk_candidate()], str(p2))
    assert ret2 == str(p2)
    with open(p2, encoding="utf-8-sig") as f:
        assert "product_id" in f.readline()


if __name__ == "__main__":
    import inspect
    import tempfile
    import traceback

    failed = total = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            total += 1
            try:
                args = (tempfile.mkdtemp(),) if inspect.signature(fn).parameters else ()
                fn(*args)
                print(f"PASS {name}")
            except Exception:
                failed += 1
                print(f"FAIL {name}")
                traceback.print_exc()
    print(f"\n{total - failed}/{total} passed")
    sys.exit(1 if failed else 0)
